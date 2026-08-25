#!/usr/bin/env python3
"""Method overview v3, matching the v6/v11 pipeline figure idiom.

Four tall blocks with dark title chips, visual elements over words,
verbs on the spine arrows, one deployment strip. Editable shapes only.
"""
from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

FONT = "Helvetica Neue"
TEXT = RGBColor(0x1F, 0x1F, 0x1E)
MUTED = RGBColor(0x6B, 0x6A, 0x63)
GRAY = RGBColor(0x8C, 0x8C, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
IN1 = RGBColor(0x4C, 0x72, 0xB0)
IN2 = RGBColor(0x5F, 0x9E, 0xA0)
IN3 = RGBColor(0x7B, 0xA7, 0xD7)
INBG = RGBColor(0xEE, 0xF2, 0xF9)
OUT1 = RGBColor(0xDD, 0x84, 0x52)
OUT2 = RGBColor(0xC4, 0x4E, 0x52)
OUTBG = RGBColor(0xFD, 0xF0, 0xE7)
GREEN = RGBColor(0x55, 0xA8, 0x68)
GREENBG = RGBColor(0xEC, 0xF5, 0xEE)
CARD = RGBColor(0xF7, 0xF6, 0xF2)
BLOCK_BG = RGBColor(0xFC, 0xFB, 0xF9)

prs = Presentation()
prs.slide_width = In(13.333)
prs.slide_height = In(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
SH = slide.shapes


def no_shadow(s):
    el = s._element.spPr
    etree.SubElement(el, qn("a:effectLst"))


def _style(tf, size, color, bold=False, italic=False, align=PP_ALIGN.CENTER):
    tf.word_wrap = True
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.italic = italic


def box(x, y, w, h, fc, ec, lw=1.2, text=None, size=10, tcolor=None,
        bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25,
        italic=False):
    s = SH.add_shape(shape, In(x), In(y), In(w), In(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    if fc is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fc
    if ec is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = ec
        s.line.width = Pt(lw)
    no_shadow(s)
    if text is not None:
        s.text_frame.text = text
        s.text_frame.margin_left = s.text_frame.margin_right = Emu(9000)
        s.text_frame.margin_top = s.text_frame.margin_bottom = Emu(4500)
        s.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _style(s.text_frame, size, tcolor or TEXT, bold=bold, italic=italic)
    return s


def label(x, y, w, h, text, size=9, color=MUTED, bold=False, italic=False,
          align=PP_ALIGN.CENTER):
    t = SH.add_textbox(In(x), In(y), In(w), In(h))
    t.text_frame.text = text
    t.text_frame.margin_left = t.text_frame.margin_right = 0
    t.text_frame.margin_top = t.text_frame.margin_bottom = 0
    _style(t.text_frame, size, color, bold=bold, italic=italic, align=align)
    return t


def conn(x0, y0, x1, y1, color, w=2.2, dash=None):
    c = SH.add_connector(MSO_CONNECTOR.STRAIGHT, In(x0), In(y0),
                         In(x1), In(y1))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    te = ln.makeelement(qn("a:tailEnd"), {"type": "triangle"})
    ln.append(te)
    if dash:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    no_shadow(c)
    return c


def pill(x, y, name, color, w=0.95, h=0.40, fs=10.5):
    return box(x, y, w, h, color, WHITE, lw=1.0, radius=0.45, text=name,
               size=fs, tcolor=WHITE, bold=True)


def spine_arrow(x, y, verb, color=GRAY, w=0.52):
    a = box(x, y, w, 0.62, color, None, shape=MSO_SHAPE.RIGHT_ARROW,
            radius=None)
    a.adjustments[0] = 0.55
    label(x - 0.24, y + 0.66, w + 0.5, 0.26, verb, size=10, color=color,
          bold=True)
    return a


def block(x, w, title, title_color=TEXT):
    box(x, 0.42, w, 5.55, BLOCK_BG, GRAY, lw=1.4, radius=0.06)
    box(x + 0.14, 0.56, w - 0.28, 0.52, title_color, None, radius=0.3,
        text=title, size=13.5, tcolor=WHITE, bold=True)


def model_chip(cx, cy, ec, name, w=1.15, h=0.78):
    box(cx - w / 2, cy - h / 2, w, h, INBG if ec != OUT1 else OUTBG, ec,
        lw=1.6)
    box(cx - 0.05, cy - h / 2 - 0.22, 0.10, 0.10, ec, None,
        shape=MSO_SHAPE.OVAL, radius=None)
    for dx in (-w * 0.18, w * 0.18):
        box(cx + dx - 0.045, cy - h * 0.10 - 0.045, 0.09, 0.09, ec, None,
            shape=MSO_SHAPE.OVAL, radius=None)
    box(cx - w * 0.14, cy + h * 0.12, w * 0.28, 0.03, ec, None,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    label(cx - w / 2 - 0.3, cy + h / 2 + 0.03, w + 0.6, 0.26, name,
          size=9.5, color=TEXT, bold=True)


# ---------------- blocks ----------------
B1, B2, B3, B4 = 0.15, 3.10, 6.05, 10.05
W1, W2, W3, W4 = 2.70, 2.70, 3.75, 3.15
block(B1, W1, "Task Demand")
block(B2, W2, "Buy Guidance", title_color=OUT1)
block(B3, W3, "Self-Explore + Verify", title_color=IN1)
block(B4, W4, "Internalize", title_color=IN2)
for gx, verb in ((2.86, "price"), (5.81, "guide"), (9.81, "learn")):
    spine_arrow(gx, 3.10, verb)

# Block 1: demand bars over atoms
label(B1 + 0.2, 1.30, W1 - 0.4, 0.3, "few-shot support set", size=10,
      color=MUTED)
for i, (nm, c, hgt) in enumerate((("plan", IN1, 1.6), ("api", IN2, 1.1),
                                  ("code", IN3, 0.7))):
    x = B1 + 0.35 + i * 0.75
    box(x, 3.55 - hgt, 0.5, hgt, c, None, radius=0.12)
    label(x - 0.1, 3.60, 0.7, 0.25, nm, size=9.5, color=TEXT)
label(B1 + 0.2, 4.05, W1 - 0.4, 0.3, "deficit δ", size=11, color=TEXT,
      bold=True)
box(B1 + 0.35, 4.45, W1 - 0.75, 0.55, CARD, GRAY, lw=1.0,
    text="unmet demand only", size=9.5, italic=True)

# Block 2: teacher chip -> context slot; crossed imitation path
model_chip(B2 + 0.75, 1.95, OUT1, "teacher")
demo = box(B2 + 1.55, 1.62, 0.85, 0.62, WHITE, OUT1, lw=1.2, text="demo",
           size=10)
conn(B2 + 1.98, 2.28, B2 + 1.98, 2.95, OUT1)
ctx = box(B2 + 1.45, 3.00, 1.05, 0.55, OUTBG, OUT1, lw=1.4, text="context",
          size=10.5, bold=True)
# crossed-out imitation target path
tgt = box(B2 + 0.30, 3.95, 1.35, 0.55, CARD, GRAY, lw=1.0, text="target",
          size=10, italic=True)
c1 = SH.add_connector(MSO_CONNECTOR.STRAIGHT, In(B2 + 0.22), In(4.58),
                      In(B2 + 1.75), In(3.88))
c1.line.color.rgb = OUT2
c1.line.width = Pt(2.6)
no_shadow(c1)
label(B2 + 0.25, 4.62, 2.2, 0.28, "never a target", size=10, color=OUT2,
      bold=True)
label(B2 + 0.2, 5.12, W2 - 0.4, 0.3, "priced per token unlocked",
      size=9.5, color=MUTED, italic=True)

# Block 3: student chip, branches to check/cross via verifier
model_chip(B3 + 0.85, 1.95, IN1, "student")
label(B3 + 0.35, 2.55, 1.9, 0.28, "unguided first", size=10, color=IN1,
      bold=True)
conn(B3 + 1.45, 2.05, B3 + 2.45, 1.75, IN1)
conn(B3 + 1.45, 2.25, B3 + 2.45, 3.35, IN1)
ok = box(B3 + 2.50, 1.50, 0.62, 0.50, GREENBG, GREEN, lw=1.6, text="✓",
         size=16, tcolor=GREEN, bold=True)
bad = box(B3 + 2.50, 3.10, 0.62, 0.50, OUTBG, OUT2, lw=1.6, text="✗",
          size=15, tcolor=OUT2, bold=True)
label(B3 + 2.20, 2.10, 1.4, 0.26, "verifier", size=10, color=TEXT,
      bold=True)
box(B3 + 0.35, 4.20, W3 - 0.7, 0.55, INBG, IN1, lw=1.2,
    text="student-written trajectories", size=10, bold=True)
label(B3 + 0.35, 4.85, W3 - 0.7, 0.28, "guided only after failure",
      size=9.5, color=MUTED, italic=True)

# Block 4: objective chips
box(B4 + 0.30, 1.60, W4 - 0.6, 0.68, GREENBG, GREEN, lw=1.4,
    text="✓  weighted log-likelihood", size=10.5, bold=True)
label(B4 + 0.30, 2.30, W4 - 0.6, 0.26, "clipped ratio → unguided policy",
      size=9.5, color=MUTED)
box(B4 + 0.30, 2.80, W4 - 0.6, 0.68, OUTBG, OUT2, lw=1.4,
    text="✓ vs ✗  from first divergence", size=10.5, bold=True)
label(B4 + 0.30, 3.50, W4 - 0.6, 0.26, "gated preference", size=9.5,
      color=MUTED)
box(B4 + 0.30, 4.15, W4 - 0.6, 0.62, GREENBG, GREEN, lw=1.4,
    text="conformal task region", size=10.5, bold=True)
label(B4 + 0.30, 4.80, W4 - 0.6, 0.26, "specialist inside, escalate outside",
      size=9.5, color=MUTED)

# re-price loop
loop = box(0.75, 6.22, 11.9, 0.5, None, GRAY, lw=1.6, radius=0.5,
           text="re-price after every update  ·  guidance fades as the "
                "student succeeds unguided  ·  stop at max U ≤ 0",
           size=10.5, tcolor=MUTED, italic=True)
conn(0.78, 6.47, 0.78, 4.30, GRAY, w=1.6)

# tagline
label(0.3, 6.95, 12.7, 0.35,
      "teacher guides · student writes · verifier decides",
      size=13, color=TEXT, bold=True)

out = "results/figs/fig_method_v3.pptx"
prs.save(out)
print("saved", out)
